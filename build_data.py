import json, re, sys
import openpyxl

XLSX = sys.argv[1]
OUT = sys.argv[2]
wb = openpyxl.load_workbook(XLSX)

def clean(v):
    if v is None:
        return ""
    return str(v).replace("\xa0", " ").strip()

# ---------- journey rows (K / V / R / C) ----------
uj = wb["User Journey"]
rows = {}
rules, clar = [], []
for r in range(1, uj.max_row + 1):
    code = clean(uj.cell(r, 2).value)
    if re.fullmatch(r"[KV]\d+", code):
        rows[code] = dict(
            code=code, actor=clean(uj.cell(r, 4).value), stage=clean(uj.cell(r, 3).value),
            activity=clean(uj.cell(r, 5).value), response=clean(uj.cell(r, 6).value),
            output=clean(uj.cell(r, 7).value),
            reqs=[c.strip() for c in clean(uj.cell(r, 8).value).split(",") if c.strip()],
            note=clean(uj.cell(r, 9).value), source=clean(uj.cell(r, 10).value))
    elif re.fullmatch(r"R\d+", code):
        rules.append(dict(code=code, text=clean(uj.cell(r, 5).value),
                          reqs=[c.strip() for c in clean(uj.cell(r, 8).value).split(",") if c.strip()],
                          note=clean(uj.cell(r, 9).value)))
    elif re.fullmatch(r"C\d+", code):
        clar.append(dict(code=code, topic=clean(uj.cell(r, 3).value), text=clean(uj.cell(r, 5).value),
                         reqs=[c.strip() for c in clean(uj.cell(r, 8).value).split(",") if c.strip()],
                         status=clean(uj.cell(r, 9).value)))
# Stage typo in source ("Konfgurasi") is kept as written in the charter.

# ---------- requirement codes ----------
pm = wb["CalculationForPM"]
reqs = {}
feature = ""
for r in range(16, 110):
    fid = clean(pm.cell(r, 3).value)
    if clean(pm.cell(r, 2).value) and not re.fullmatch(r"(Fitur|Total.*)", clean(pm.cell(r, 2).value)):
        feature = clean(pm.cell(r, 2).value)
    if re.fullmatch(r"(BE|FE|DS|DP)-\d+[A-Z]", fid):
        reqs[fid] = dict(code=fid, feature=feature, user=clean(pm.cell(r, 4).value),
                         desc=clean(pm.cell(r, 5).value), refs=clean(pm.cell(r, 9).value),
                         advice=clean(pm.cell(r, 10).value))

SUMMARY = {
 "BE-01A": "Login/logout, menjaga sesi, menolak akses yang tidak sah.",
 "BE-01B": "Role dan hak akses fungsi/Project untuk tiap pengguna.",
 "BE-01C": "Kelola akun, undangan/registrasi, verifikasi email, reset password.",
 "BE-02A": "Buat dan ubah seluruh field Task tanpa menimpa hasil yang sudah ada.",
 "BE-03A": "Task One-Time dan recurring otomatis mengikuti pola assignment sebelumnya.",
 "BE-03B": "Dua metode pembagian: rentang shift/timestamp atau jumlah data (Validation Unit).",
 "BE-04A": "Task Pool per Project: daftar dan detail Task keempat media.",
 "BE-04B": "Progress, sisa workload, overdue, dan completion berbasis Validation Unit.",
 "BE-04C": "Notifikasi 60 dan 30 menit sebelum due date, serta peringatan overdue/gagal.",
 "BE-05A": "Deteksi data bertambah/berkurang, beri notifikasi, buat Task tambahan.",
 "BE-06A": "Assignment Validator dan Start/Continue, terbuka setelah lulus pretest.",
 "BE-07A": "Stream Social Media lengkap dengan konteks dan keputusan per Object.",
 "BE-08A": "Artikel Online News sesuai Clipping dan timeframe Task.",
 "BE-09A": "Stream Print dan Broadcast dengan metadata dan attachment.",
 "BE-10A": "Tetapkan sentimen Positive/Neutral/Negative dan tandai Validated.",
 "BE-10B": "Pilih atau lepas Label existing; Validator tidak membuat master Label.",
 "BE-10C": "Terjemahkan teks berbahasa asing ke Bahasa Indonesia.",
 "BE-10D": "Ignore dengan alasan dan Restore tanpa mengubah data lain.",
 "BE-11A": "Pencarian dan kombinasi filter, termasuk include/exclude Label.",
 "BE-12A": "Kelompokkan exact duplicate: satu representative beserta member.",
 "BE-12B": "Kelompokkan Retweet berdasarkan parent yang sama.",
 "BE-12C": "Bulk action ke item terpilih dengan jumlah item terdampak.",
 "BE-12D": "Saran hasil validasi Retweet dari Task/timeframe sebelumnya.",
 "BE-13A": "Draft tersimpan otomatis dan bisa dilanjutkan.",
 "BE-13B": "Satu Submit per Assignment; Completed setelah semua hasil masuk dashboard.",
 "BE-13C": "Kirim hasil Social Media ke record dan Object yang tepat.",
 "BE-13D": "Kirim hasil Online Media ke Clipping dan record yang tepat.",
 "BE-13E": "Kirim hasil Print/Broadcast ke record sumber yang tepat.",
 "BE-14A": "Metrik performa Validator: productivity, completion, akurasi, ignore rate.",
 "BE-15A": "Pretest wajib 10 soal, boleh diulang sampai semua benar.",
 "BE-16A": "Activity Center untuk memantau proses perpindahan data.",
 "FE-01A": "Halaman login/logout, session expired, dan access denied.",
 "FE-01B": "Halaman undangan, registrasi, verifikasi, dan reset password.",
 "FE-01C": "Halaman kelola akun, status akun, Role, dan otorisasi.",
 "FE-02A": "Form buat/edit Task dengan seluruh field.",
 "FE-03A": "Pilih metode pembagian, Validator, jadwal recurring, dan override.",
 "FE-04A": "Monitoring Task: progress, overdue, notifikasi, akses modify/reassign.",
 "FE-05A": "Notifikasi dan detail perubahan data, plus aksi buat Task tambahan.",
 "FE-06A": "Daftar Assignment Validator; Start/Continue setelah lulus pretest.",
 "FE-07A": "Workspace Social Media: sentimen, Label, Ignore/Restore, Translate.",
 "FE-08A": "Workspace Online Media dengan aksi validasi yang sama.",
 "FE-09A": "Workspace Print/Broadcast dengan aksi validasi yang sama.",
 "FE-11A": "Pencarian dan filter; pilih item dari hasil filter.",
 "FE-12A": "Tampilan duplicate/Retweet: representative, member, konfirmasi bulk.",
 "FE-13A": "Status Draft, konfirmasi Submit, dan status pengiriman.",
 "FE-14A": "Halaman profil performa Validator.",
 "FE-15A": "Coordinator menyiapkan 10 soal pretest dan kunci jawabannya.",
 "FE-15B": "Validator mengerjakan pretest sampai lulus.",
 "FE-16A": "Panel Activity Center untuk semua pengguna.",
 "FE-17A": "Kualitas UI: desktop, konsisten dengan NoLimit, state jelas.",
 "DS-01A": "Desain akses, akun, Role, dan otorisasi.",
 "DS-02A": "Desain form Task dan review perubahan.",
 "DS-03A": "Desain pembagian kerja, recurring, dan reassignment.",
 "DS-04A": "Desain monitoring Task, notifikasi, dan status.",
 "DS-05A": "Desain notifikasi dan detail perubahan data sumber.",
 "DS-06A": "Desain pekerjaan Validator dan kondisi terkunci pretest.",
 "DS-07A": "Desain workspace Social Media beserta seluruh interaksinya.",
 "DS-08A": "Desain workspace Online Media.",
 "DS-09A": "Desain workspace Print/Broadcast.",
 "DS-11A": "Desain pencarian dan filter.",
 "DS-12A": "Desain duplicate/Retweet dan bulk action.",
 "DS-13A": "Desain Draft, Submit, dan status pengiriman.",
 "DS-14A": "Desain detail performa Validator.",
 "DS-15A": "Desain pretest untuk Coordinator dan Validator.",
 "DS-16A": "Desain Activity Center sesuai referensi.",
 "DS-17A": "Komponen dan state UI yang konsisten.",
 "DS-18A": "Peta alur, aturan scope, dan status.",
 "DS-18B": "Navigasi dan hubungan antarhalaman.",
 "DP-19A": "Environment dan staging tersedia.",
 "DP-19B": "Akses aplikasi lewat koneksi aman.",
 "DP-19C": "Berjalan konsisten di development, staging, production.",
 "DP-19D": "Build dan deployment terkontrol serta tertelusur.",
 "DP-19E": "Pemeriksaan otomatis sebelum rilis.",
 "DP-19F": "Kesiapan production, smoke test, dan rollback.",
 "DP-20A": "Penelusuran proses, health layanan, dan error.",
}
missing = [c for c in reqs if c not in SUMMARY]
assert not missing, missing
for c, s in SUMMARY.items():
    assert c in reqs, c
    reqs[c]["summary"] = s

# ---------- lofi screens (PDF page -> screen) ----------
S = lambda p, t, a, g: dict(page=p, title=t, actor=a, group=g)
screens = [
 S(1, "Referensi NoLimit existing: Manage Social Profile", "Referensi", "Referensi dashboard existing"),
 S(2, "Referensi NoLimit existing: Activity Center", "Referensi", "Referensi dashboard existing"),
 S(3, "Onboarding: Welcome", "Validator", "Onboarding & pre-test"),
 S(4, "Onboarding: Task Guideline", "Validator", "Onboarding & pre-test"),
 S(5, "Onboarding: intro pre-test", "Validator", "Onboarding & pre-test"),
 S(6, "Onboarding: mengerjakan pre-test", "Validator", "Onboarding & pre-test"),
 S(7, "Dashboard Coordinator", "Coordinator", "Monitoring Coordinator"),
 S(8, "Validator's Profile (performa)", "Coordinator", "Monitoring Coordinator"),
 S(9, "Password Updated", "Semua", "Akses & akun"),
 S(10, "Reset Password", "Semua", "Akses & akun"),
 S(11, "Check Your Email", "Semua", "Akses & akun"),
 S(12, "Onboarding: pre-test di bawah pass mark", "Validator", "Onboarding & pre-test"),
 S(13, "Onboarding: pre-test lulus", "Validator", "Onboarding & pre-test"),
 S(14, "Password Recovery", "Semua", "Akses & akun"),
 S(15, "Login", "Semua", "Akses & akun"),
 S(16, "Home: pre-test di bawah pass mark", "Validator", "Onboarding & pre-test"),
 S(17, "Home: pre-test lulus", "Validator", "Onboarding & pre-test"),
 S(18, "Task & Assignment: Task Pool", "Coordinator", "Task & Assignment"),
 S(19, "Task & Assignment: daftar Tasks", "Coordinator", "Task & Assignment"),
 S(20, "Session Expired", "Semua", "Akses & akun"),
 S(21, "Access Denied", "Semua", "Akses & akun"),
 S(22, "Link Expired", "Semua", "Akses & akun"),
 S(23, "Create Account (dari undangan)", "Semua", "Akses & akun"),
 S(24, "Home Validator: Today's work", "Validator", "Home & notifikasi Validator"),
 S(25, "Create Task", "Coordinator", "Task & Assignment"),
 S(26, "Create Task: memilih Object", "Coordinator", "Task & Assignment"),
 S(27, "Edit Task", "Coordinator", "Task & Assignment"),
 S(28, "Waiting for Approval", "Semua", "Akses & akun"),
 S(29, "New assignment: Task Guideline", "Validator", "Onboarding & pre-test"),
 S(30, "New assignment: intro pre-test", "Validator", "Onboarding & pre-test"),
 S(31, "New assignment: mengerjakan pre-test", "Validator", "Onboarding & pre-test"),
 S(32, "New assignment: pre-test di bawah pass mark", "Validator", "Onboarding & pre-test"),
 S(33, "Kunci jawaban pre-test (Coordinator)", "Coordinator", "Monitoring Coordinator"),
 S(34, "Hasil pre-test/post-test vs kunci", "Coordinator", "Monitoring Coordinator"),
 S(35, "New assignment: pre-test lulus", "Validator", "Onboarding & pre-test"),
 S(36, "Manage Account: approval akun", "Coordinator", "Akses & akun"),
 S(37, "Workspace Social Media (halaman penuh)", "Validator", "Workspace Social Media"),
 S(38, "Social Media: pilih semua di halaman", "Validator", "Workspace Social Media"),
 S(39, "Social Media: pilih seluruh 1.500 stream", "Validator", "Workspace Social Media"),
 S(40, "Social Media: bulk sentimen + terapkan ke RT", "Validator", "Workspace Social Media"),
 S(41, "Notifikasi Coordinator", "Coordinator", "Monitoring Coordinator"),
 S(42, "Social Media: pilih emotion", "Validator", "Workspace Social Media"),
 S(43, "Social Media: pilih timeframe", "Validator", "Workspace Social Media"),
 S(44, "Social Media: tab Draft", "Validator", "Workspace Social Media"),
 S(45, "Social Media: View Ignored + Restore", "Validator", "Workspace Social Media"),
 S(46, "Social Media: Ignore dengan alasan", "Validator", "Workspace Social Media"),
 S(47, "Social Media: konfirmasi POST", "Validator", "Workspace Social Media"),
 S(48, "Social Media: 38 copies (exact duplicate)", "Validator", "Workspace Social Media"),
 S(49, "Social Media: Advance search", "Validator", "Pencarian & filter"),
 S(50, "Social Media: Advance search (lengkap)", "Validator", "Pencarian & filter"),
 S(51, "Social Media: Manage Social Profile", "Validator", "Pencarian & filter"),
 S(52, "Activity Center (lofi)", "Semua", "Workspace Social Media"),
 S(53, "Workspace Online Media", "Validator", "Workspace Online/Print/Broadcast"),
 S(54, "Workspace Print Media", "Validator", "Workspace Online/Print/Broadcast"),
 S(55, "Workspace Broadcast Media", "Validator", "Workspace Online/Print/Broadcast"),
 S(56, "Online Media: Filter", "Validator", "Pencarian & filter"),
 S(57, "Print Media: Filter", "Validator", "Pencarian & filter"),
 S(58, "Broadcast Media: Filter", "Validator", "Pencarian & filter"),
 S(59, "Online Media: Filter (lengkap)", "Validator", "Pencarian & filter"),
 S(60, "Print Media: Filter (lengkap)", "Validator", "Pencarian & filter"),
 S(61, "Broadcast Media: Filter (lengkap)", "Validator", "Pencarian & filter"),
 S(62, "Offboarding: intro post-test", "Validator", "Offboarding & post-test"),
 S(63, "Offboarding: mengerjakan post-test", "Validator", "Offboarding & post-test"),
 S(64, "Offboarding: post-test selesai", "Validator", "Offboarding & post-test"),
 S(65, "Notifikasi Validator", "Validator", "Home & notifikasi Validator"),
 S(66, "Home Validator + panel notifikasi", "Validator", "Home & notifikasi Validator"),
]
assert [s["page"] for s in screens] == list(range(1, 67))

# ---------- 4 journeys: step -> screens, lofi status, note ----------
# status: ada | sebagian | belum
PASS_NOTE = ("Lofi memakai pass mark 9/10 dan \"one attempt\", sedangkan C12/BE-15A/FE-15A "
             "menetapkan 10 soal harus benar semua dan boleh diulang. Perlu diselaraskan.")
STEP = {
 "K2":  ([15, 14, 11, 10, 9, 23, 22, 28, 20, 21, 36, 7], "ada", ""),
 "V2":  ([15, 23, 28, 22, 20, 21], "ada", ""),
 "V3":  ([24, 65, 66], "ada", ""),
 "V4":  ([24, 29, 37], "ada", ""),
 "V5":  ([29, 4, 24], "ada", ""),
 "V6":  ([3, 4, 5, 6, 12, 13, 29, 30, 31, 32, 35, 16, 17], "ada", PASS_NOTE),
 "K3":  ([18, 19, 25], "ada", ""),
 "K4":  ([25], "ada", "Project dipilih pada Data scope; Project aktif juga bisa diganti lewat tombol Switch di sidebar."),
 "K5":  ([25], "ada", "Tipe task diwakili toggle Recurring (off = One-Time)."),
 "K6":  ([25], "ada", "Isian Repeat every, Create at, Ends."),
 "K7":  ([25], "ada", "Pilihan By stream count (workload) / By timeframe (shift) di bagian Validators."),
 "K8":  ([25], "ada", "Data date range From - To."),
 "K9":  ([25], "ada", "Pilihan Source: Social, Online, Print, Broadcast Media."),
 "K10": ([25, 26], "sebagian", "Group, Object, dan Label sudah ada untuk Social Media. Pilihan Clipping untuk Online/Print/Broadcast belum terlihat."),
 "K11": ([25], "belum", "Belum terlihat pilihan action yang diizinkan (Validation, Ignore/Restore, Labeling) pada form Create Task."),
 "K12": ([25], "sebagian", "Validator dan pembagian jumlah stream sudah ada. Info workload/availability tiap Validator dan urutan rekomendasi belum terlihat."),
 "K13": ([25], "sebagian", "Tab By timeframe tersedia, tetapi isi konfigurasi shift (jumlah shift dan Validator per shift) belum terlihat."),
 "K14": ([25], "ada", ""),
 "K15": ([25], "ada", ""),
 "K16": ([25], "ada", "Panel Summary dan estimasi jumlah stream (57,611 streams match)."),
 "K17": ([25], "ada", "Tombol Create & assign dan Save as draft."),
 "K19": ([19, 27], "ada", ""),
 "K20": ([27], "ada", ""),
 "K21": ([27], "sebagian", "Summary hanya menampilkan konfigurasi terbaru; perbandingan sebelum vs sesudah belum terlihat."),
 "K22": ([27], "ada", "Tombol Save changes."),
 "V7":  ([37, 38, 39, 40, 42, 43, 53, 54, 55, 49, 50, 51, 56, 57, 58, 59, 60, 61], "ada", "Translate (R3, BE-10C) belum terlihat di workspace."),
 "V8":  ([46, 45, 44], "ada", ""),
 "V9":  ([37, 53], "belum", "Pilih/lepas Label existing di workspace belum terlihat."),
 "V10": ([37, 48, 40], "ada", "Quote/reply, jumlah copies, dan opsi menerapkan sentimen ke semua RT."),
 "V11": ([37, 24], "ada", "Progress bar unit di bagian bawah workspace dan kartu di Home."),
 "V12": ([44, 37, 62], "ada", ""),
 "V13": ([47, 52, 62, 63, 64], "ada", "Lofi menambahkan post-test setelah submit; post-test belum tercantum di project charter."),
 "K18": ([19, 18, 7], "ada", ""),
 "K23": ([7, 19], "ada", ""),
 "K24": ([7, 41], "belum", "Belum terlihat tampilan penambahan/pengurangan data pada timeframe Task. Panel Sync dan notifikasi Sync failed adalah status sinkronisasi, bukan delta data."),
 "K25": ([], "belum", "Detail data tambahan yang belum masuk Task belum terlihat."),
 "K26": ([], "belum", "Aksi membuat Task tambahan dari data baru belum terlihat."),
 "K27": ([19, 7, 41], "sebagian", "Status Overdue dan due date sudah ada. Notifikasi 60 dan 30 menit sebelum due date belum terlihat."),
 "K28": ([], "belum", "Progress per shift belum terlihat."),
 "K29": ([7, 8, 33, 34], "ada", ""),
 "K30": ([41, 7], "ada", "Notifikasi Coordinator, Priority queue, dan panel Sync dengan Retry."),
 "K31": ([19, 27], "sebagian", "Edit Task tersedia (termasuk ganti Validator). Aksi khusus Reassign, Adjust Shift, dan Create Additional Task belum terlihat."),
 "K32": ([19, 41], "ada", "Status Submitted/Validated di daftar Tasks dan notifikasi submit."),
}
journeys = [
 dict(id="J1", title="Akses & Onboarding", actor="Coordinator + Validator",
      blurb="Masuk panel, lalu Validator melihat task, membaca guideline, dan lulus pre-test sebelum mulai.",
      steps=["K2", "V2", "V3", "V4", "V5", "V6"]),
 dict(id="J2", title="Membuat & Mengubah Task", actor="Coordinator",
      blurb="Coordinator menyusun task dari project, scope, jadwal, sampai pembagian Validator, lalu mengubahnya bila perlu.",
      steps=["K3", "K4", "K5", "K6", "K7", "K8", "K9", "K10", "K11", "K12", "K13", "K14", "K15", "K16", "K17", "K19", "K20", "K21", "K22"]),
 dict(id="J3", title="Validasi Stream sampai POST", actor="Validator",
      blurb="Validator memvalidasi sentimen, Ignore/Restore, Label, duplicate/Retweet, lalu mengirim hasil.",
      steps=["V7", "V8", "V9", "V10", "V11", "V12", "V13"]),
 dict(id="J4", title="Monitoring & Kontrol", actor="Coordinator",
      blurb="Coordinator memantau task, progress, data baru, SLA, performa Validator, dan menindaklanjuti masalah.",
      steps=["K18", "K23", "K24", "K25", "K26", "K27", "K28", "K29", "K30", "K31", "K32"]),
]
all_steps = [s for j in journeys for s in j["steps"]]
assert sorted(all_steps) == sorted(rows), set(rows) ^ set(all_steps)
assert len(all_steps) == len(set(all_steps))
# steps whose lofi conflicts with or goes beyond the charter
FLAG = {"V6", "V13"}
for code, (pages, status, note) in STEP.items():
    rows[code].update(screens=pages, lofi=status, lofiNote=note, lofiFlag=code in FLAG)
for j in journeys:
    for s in j["steps"]:
        rows[s]["journey"] = j["id"]

# codes present in every journey row = foundation
foundation = sorted(set.intersection(*[set(r["reqs"]) for r in rows.values()]))
unknown = {c for r in rows.values() for c in r["reqs"]} - set(reqs)
assert not unknown, unknown

data = dict(rows=rows, journeys=journeys, reqs=reqs, screens=screens, rules=rules,
            clar=clar, foundation=foundation, generated="Project charter 22 Oktober 2026 · Figma lofi iterasi 4")
with open(OUT, "w") as f:
    f.write("window.DATA = " + json.dumps(data, ensure_ascii=False) + ";\n")
print("rows", len(rows), "reqs", len(reqs), "foundation", foundation, "rules", len(rules), "clar", len(clar))
